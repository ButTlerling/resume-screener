import os
import re
import json
import time
import pandas as pd
import streamlit as st
import pdfplumber
import docx
import anthropic
from openai import OpenAI
from google import genai as google_genai
from dotenv import load_dotenv

load_dotenv()

# ── Constants ──────────────────────────────────────────────────────────────────
CLAUDE_MODEL = "claude-sonnet-4-6"
GPT_MODEL = "gpt-4o"
GPT_54_MODEL = "gpt-5.4"
GEMINI_FLASH_MODEL = "gemini-2.5-flash"
GEMINI_PRO_MODEL = "gemini-2.5-pro"
GEMINI_31_PRO_MODEL = "gemini-3.1-pro-preview"
MAX_RESUME_CHARS = 24_000
MAX_TOKENS_RESPONSE = 1024

MODEL_OPTIONS = {
    "GPT-5.4": "openai_54",
    "GPT-4o": "openai",
    "Claude (claude-sonnet-4-6)": "claude",
    "Gemini 3.1 Pro": "gemini_31_pro",
    "Gemini 2.5 Pro": "gemini_pro",
    "Gemini 2.5 Flash": "gemini_flash",
}

API_KEY_MAP = {
    "claude": "ANTHROPIC_API_KEY",
    "openai": "OPENAI_API_KEY",
    "openai_54": "OPENAI_API_KEY",
    "gemini_31_pro": "GOOGLE_API_KEY",
    "gemini_pro": "GOOGLE_API_KEY",
    "gemini_flash": "GOOGLE_API_KEY",
}

SYSTEM_PROMPT = """You are an expert technical recruiter and hiring manager.
You will receive a job description and a resume, and you must evaluate how well
the candidate matches the role.

You MUST respond with valid JSON only. No prose before or after the JSON.
Use exactly this schema:

{
  "candidate_name": "<full name of the candidate, or 'Unknown' if not found>",
  "match_score": <integer 1-10>,
  "key_strengths": [<string>, ...],
  "key_gaps": [<string>, ...],
  "summary": "<one sentence summary of the candidate>",
  "university": "<university name, or 'Not listed' if absent>",
  "major": "<degree and field of study, e.g. 'B.S. Computer Science', or 'Not listed' if absent>",
  "skills": {
    "coding": [<programming languages explicitly listed, e.g. "Python", "Java", "C++">],
    "libraries": [<frameworks/libraries listed, e.g. "NumPy", "TensorFlow", "React">],
    "electrical_eng": [<EE-related skills listed, e.g. "MATLAB", "PCB Design", "Circuit Analysis"; empty list if none>],
    "other": [<other technical skills not in above categories, e.g. "SQL", "Excel", "CAD"; empty list if none>]
  },
  "internships": [<"Company — Role (Year)" for each internship, max 5 items, empty list if none>],
  "projects": [<brief project title or one-line description, max 5 items, empty list if none>],
  "extracurriculars": [<clubs, sports, volunteer work, leadership, hobbies, max 5 items, empty list if none>]
}

Scoring rubric — blended skills + character system (total 10 points):

=== PART A: SKILLS SCORE (0-7 points) ===

STEP 1 — Base score from REQUIRED skills (1-4 points):
   4: Meets 80%+ of required skills
   3: Meets 60-79% of required skills
   2: Meets 40-59% of required skills
   1: Meets <40% of required skills, completely different field, or unreadable resume

STEP 2 — Bonus from PREFERRED/nice-to-have skills (only if base >= 2):
   Each preferred skill present adds +1, up to a maximum of +3 bonus points.
   Preferred skills have NO effect if base score is 1.

Skills score = base + preferred bonus (max 7).

=== PART B: CHARACTER SCORE (0-3 points) ===

Look for evidence of personality, initiative, and moldability beyond academics:

   3 points: Two or more STRONG signals — military service, sustained volunteer work (semester+),
             elected club officer/team captain, meaningful personal project driven by passion
   2 points: One strong signal OR two minor signals — club member, one-time volunteer,
             hackathon participant, part-time job during school, casual hobby mentioned
   1 point:  One minor mention — e.g. "enjoys hiking", one club listed with no role
   0 points: Nothing beyond academics and internships; no extracurriculars or personal interests

Strong signals show the candidate is coachable, disciplined, and eager to grow.
Military experience always counts as a strong signal regardless of field.

=== STEP 3 — FINAL SCORE ===

Final score = Skills score + Character score, capped at 10.

Examples:
   Skills base=4 (strong), preferred +3, character=3 (veteran + volunteer) = 7+3 = 10
   Skills base=4 (strong), preferred +2, character=0 (no activities)       = 6+0 = 6
   Skills base=3 (partial), preferred +1, character=3 (club officer + hobby) = 4+3 = 7
   Skills base=2 (weak),    preferred +0, character=3 (military veteran)    = 2+3 = 5
   Skills base=1 (poor),    preferred +0, character=2 (volunteer)           = 1+2 = 3

Apply mechanically. Do not adjust scores based on subjective impressions beyond the rubric.
"""

USER_TEMPLATE = """JOB DESCRIPTION:
{job_description}

---

RESUME (filename: {filename}):
{resume_text}

---

Evaluate this resume against the job description and respond with JSON only."""


# ── Helper functions ───────────────────────────────────────────────────────────

def extract_text_from_pdf(uploaded_file) -> str | None:
    """Extract full text from a PDF UploadedFile. Returns None on failure."""
    try:
        text_parts = []
        with pdfplumber.open(uploaded_file) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text.strip())
        if not text_parts:
            return None
        full_text = "\n\n--- PAGE BREAK ---\n\n".join(text_parts)
        return full_text[:MAX_RESUME_CHARS]
    except Exception:
        return None


def extract_text_from_docx(uploaded_file) -> str | None:
    """Extract full text from a DOCX UploadedFile. Returns None on failure."""
    try:
        doc = docx.Document(uploaded_file)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        if not paragraphs:
            return None
        return "\n\n".join(paragraphs)[:MAX_RESUME_CHARS]
    except Exception:
        return None


def extract_text(uploaded_file) -> str | None:
    """Route to correct extractor based on file extension."""
    name = uploaded_file.name.lower()
    if name.endswith(".pdf"):
        return extract_text_from_pdf(uploaded_file)
    if name.endswith(".docx"):
        return extract_text_from_docx(uploaded_file)
    return None


def _page_has_contact_header(page_text: str) -> bool:
    """Check if a page starts with typical resume contact info.

    Counts distinct contact signals in the first 500 characters and requires
    at least 2.  This catches resumes that use hyperlink placeholders
    (e.g. 'Email | LinkedIn') instead of showing the raw address.
    """
    header = page_text[:500]
    signals = 0
    if re.search(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', header):
        signals += 1  # raw email address
    if re.search(r'[\+]?[\d\-\(\)\s]{7,}', header):
        signals += 1  # phone number
    if re.search(r'\bEmail\b', header, re.IGNORECASE):
        signals += 1  # "Email" as hyperlink placeholder
    if re.search(r'linkedin', header, re.IGNORECASE):
        signals += 1  # LinkedIn profile link or word
    if re.search(r'github', header, re.IGNORECASE):
        signals += 1  # GitHub profile link or word
    return signals >= 2


def _extract_contact_id(page_text: str) -> tuple[str | None, str | None]:
    """Extract the first email and first phone from a page header.

    Returns (email_or_None, phone_digits_or_None) for identity matching.
    """
    header = page_text[:500]
    email_match = re.search(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', header)
    phone_match = re.search(r'[\+]?[\d\-\(\)\s]{7,}', header)
    email = email_match.group(0).lower() if email_match else None
    phone = re.sub(r'\D', '', phone_match.group(0)) if phone_match else None
    return email, phone


def split_combined_pdf(uploaded_file) -> list[tuple[str, str]]:
    """Split a combined PDF into individual resume texts by detecting contact-info boundaries.

    Each page whose header contains contact signals is treated as the start of
    a new resume.  Pages without are continuations.  After initial splitting,
    groups that share the same name AND (email or phone) are merged back
    together (handles resumes that repeat their header on every page).

    Returns a list of (label, text) tuples.
    """
    base_name = uploaded_file.name
    try:
        with pdfplumber.open(uploaded_file) as pdf:
            pages_text: list[str | None] = []
            for page in pdf.pages:
                raw = page.extract_text()
                pages_text.append(raw.strip() if raw else None)
    except Exception:
        return []

    # Group pages into resume chunks based on contact-info detection
    groups: list[list[int]] = []  # each item is a list of page indices
    for i, text in enumerate(pages_text):
        if text is None:
            continue  # skip blank pages
        if i == 0 or _page_has_contact_header(text):
            groups.append([i])  # start a new resume
        else:
            if groups:
                groups[-1].append(i)
            else:
                groups.append([i])  # safety fallback

    # Merge groups that share the same identity (name + email or phone)
    def _group_identity(page_indices: list[int]) -> tuple[str, str | None, str | None]:
        first_text = pages_text[page_indices[0]] or ""
        name = first_text.split("\n")[0].strip().lower()
        email, phone = _extract_contact_id(first_text)
        return name, email, phone

    merged: list[list[int]] = []
    for group in groups:
        name, email, phone = _group_identity(group)
        # Try to find an existing merged group with matching identity
        found = False
        for existing in merged:
            ename, eemail, ephone = _group_identity(existing)
            if name and name == ename:
                if (email and email == eemail) or (phone and phone == ephone):
                    existing.extend(group)
                    found = True
                    break
        if not found:
            merged.append(group)

    entries: list[tuple[str, str]] = []
    for resume_num, page_indices in enumerate(merged, start=1):
        text_parts = [pages_text[idx] for idx in page_indices if pages_text[idx]]
        if not text_parts:
            continue

        full_text = "\n\n--- PAGE BREAK ---\n\n".join(text_parts)
        full_text = full_text[:MAX_RESUME_CHARS]

        first = page_indices[0] + 1
        last = page_indices[-1] + 1
        if first == last:
            label = f"{base_name} — Resume #{resume_num} (page {first})"
        else:
            label = f"{base_name} — Resume #{resume_num} (pages {first}-{last})"

        entries.append((label, full_text))

    return entries


def _parse_json(raw: str) -> dict:
    """Strip markdown fences and parse JSON."""
    raw = raw.strip()
    if raw.startswith("```"):
        parts = raw.split("```")
        raw = parts[1] if len(parts) > 1 else raw
        if raw.startswith("json"):
            raw = raw[4:]
    return json.loads(raw.strip())


def analyze_resume_claude(
    client: anthropic.Anthropic,
    job_description: str,
    resume_text: str,
    filename: str,
) -> dict | None:
    prompt = USER_TEMPLATE.format(
        job_description=job_description,
        filename=filename,
        resume_text=resume_text,
    )
    for attempt in range(2):
        try:
            message = client.messages.create(
                model=CLAUDE_MODEL,
                max_tokens=MAX_TOKENS_RESPONSE,
                temperature=0,
                system=SYSTEM_PROMPT,
                messages=[{"role": "user", "content": prompt}],
            )
            return _parse_json(message.content[0].text)
        except json.JSONDecodeError:
            if attempt == 0:
                continue
            return None
        except anthropic.RateLimitError:
            time.sleep(10)
            if attempt == 0:
                continue
            return None
        except Exception as e:
            return {"_error": str(e)}
    return None


def analyze_resume_openai(
    client: OpenAI,
    job_description: str,
    resume_text: str,
    filename: str,
    model: str = GPT_MODEL,
) -> dict | None:
    prompt = USER_TEMPLATE.format(
        job_description=job_description,
        filename=filename,
        resume_text=resume_text,
    )
    for attempt in range(2):
        try:
            params = dict(
                model=model,
                temperature=0,
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": prompt},
                ],
            )
            if model == GPT_54_MODEL:
                params["max_completion_tokens"] = MAX_TOKENS_RESPONSE
            else:
                params["max_tokens"] = MAX_TOKENS_RESPONSE
            response = client.chat.completions.create(**params)
            return _parse_json(response.choices[0].message.content)
        except json.JSONDecodeError:
            if attempt == 0:
                continue
            return None
        except Exception as e:
            return {"_error": str(e)}
    return None


def analyze_resume_gemini(
    client,
    job_description: str,
    resume_text: str,
    filename: str,
    model: str = GEMINI_FLASH_MODEL,
) -> dict | None:
    prompt = SYSTEM_PROMPT + "\n\n" + USER_TEMPLATE.format(
        job_description=job_description,
        filename=filename,
        resume_text=resume_text,
    )
    for attempt in range(2):
        try:
            response = client.models.generate_content(
                model=model,
                contents=prompt,
            )
            return _parse_json(response.text)
        except json.JSONDecodeError:
            if attempt == 0:
                continue
            return None
        except Exception as e:
            return {"_error": str(e)}
    return None


def analyze_resume(client, provider: str, job_description: str, resume_text: str, filename: str) -> dict | None:
    if provider == "claude":
        return analyze_resume_claude(client, job_description, resume_text, filename)
    if provider == "openai_54":
        return analyze_resume_openai(client, job_description, resume_text, filename, model=GPT_54_MODEL)
    if provider == "gemini_31_pro":
        return analyze_resume_gemini(client, job_description, resume_text, filename, model=GEMINI_31_PRO_MODEL)
    if provider == "gemini_pro":
        return analyze_resume_gemini(client, job_description, resume_text, filename, model=GEMINI_PRO_MODEL)
    if provider == "gemini_flash":
        return analyze_resume_gemini(client, job_description, resume_text, filename, model=GEMINI_FLASH_MODEL)
    return analyze_resume_openai(client, job_description, resume_text, filename)



def build_dataframe(results: list[dict]) -> pd.DataFrame:
    rows = []
    for rank, r in enumerate(results, start=1):
        rows.append(
            {
                "Rank": rank,
                "Name": r.get("candidate_name", "Unknown"),
                "Score (1-10)": r["match_score"],
                "University": r.get("university", "Not listed"),
                "Major": r.get("major", "Not listed"),
"Key Strengths": " | ".join(r.get("key_strengths", [])),
                "Key Gaps": " | ".join(r.get("key_gaps", [])),
                "Summary": r.get("summary", ""),
                "Filename": r["filename"],
            }
        )
    return pd.DataFrame(rows)


# ── Streamlit UI ───────────────────────────────────────────────────────────────

_STOPWORDS = {
    'the','a','an','and','or','but','in','on','at','to','for','of','with','by',
    'from','up','about','into','is','are','was','were','be','been','have','has',
    'had','do','does','did','will','would','could','should','may','might','must',
    'can','this','that','these','those','we','our','you','your','it','its','as',
    'if','not','no','so','than','then','their','they','them','what','which','who',
    'when','where','how','all','each','both','few','more','most','other','some',
    'such','only','same','also','well','just','very','new','work','use','using',
    'used','including','include','experience','strong','knowledge','ability',
    'skills','skill','role','candidate','team','position','responsibilities',
}

def _jd_keywords(jd: str) -> set[str]:
    """Extract meaningful lowercase keywords from the job description."""
    words = re.findall(r'[a-zA-Z][a-zA-Z0-9+#.]*', jd.lower())
    return {w for w in words if w not in _STOPWORDS and len(w) > 2}

def _skill_matches(skill: str, jd_text: str) -> bool:
    """True only when the skill name itself (as a whole word/phrase) appears in the JD."""
    jd_lower = jd_text.lower()
    skill_lower = re.sub(r'[-/]', ' ', skill.lower()).strip()
    # Phrase match: the full skill name is a substring surrounded by word boundaries
    pattern = r'(?<![a-zA-Z0-9])' + re.escape(skill_lower) + r'(?![a-zA-Z0-9])'
    return bool(re.search(pattern, jd_lower))

def _text_matches(text: str, jd_kw: set[str]) -> bool:
    words = re.findall(r'[a-zA-Z][a-zA-Z0-9+#.]*', text.lower())
    hits = sum(1 for w in words if w in jd_kw and len(w) >= 4)
    return hits >= 3


@st.dialog("Candidate Details", width="large")
def show_candidate_dialog(r: dict, jd_kw: set[str], jd_text: str):
    # Force dialog to fill the screen width
    st.markdown(
        """
<style>
div[data-testid="stDialog"] > div > div[role="dialog"] {
    width: 92vw !important;
    max-width: 92vw !important;
}
</style>
""",
        unsafe_allow_html=True,
    )

    name = r.get("candidate_name", "Unknown")
    score = r.get("match_score", "N/A")
    skills = r.get("skills", {})

    # ── Header ──
    score_color = "#16a34a" if score >= 8 else "#d97706" if score >= 5 else "#dc2626"
    st.markdown(
        f"<h2 style='margin:0 0 0.2rem 0;'>{name}"
        f"&ensp;<span style='color:{score_color}; font-size:1.1rem; font-weight:600;'>{score}/10</span></h2>",
        unsafe_allow_html=True,
    )
    st.markdown(
        f"<p style='color:#6b7280; margin:0 0 0.3rem 0;'>{r.get('summary', '')}</p>"
        f"<p style='font-size:0.85rem; color:#9ca3af; margin:0;'>"
        f"{r.get('university', 'Not listed')} &nbsp;·&nbsp; {r.get('major', 'Not listed')}</p>",
        unsafe_allow_html=True,
    )
    st.divider()

    # ── Skills row (full width, 4 inline blocks) ──
    coding = skills.get("coding", [])
    libs   = skills.get("libraries", [])
    ee     = skills.get("electrical_eng", [])
    other  = skills.get("other", [])

    def _tags(items: list[str]) -> str:
        if not items:
            return "<span style='color:#9ca3af; font-size:0.82rem;'>None listed</span>"
        parts = []
        for t in items:
            if _skill_matches(t, jd_text):
                bg, fg = "#dbeafe", "#1d4ed8"   # calm blue — JD match
            else:
                bg, fg = "#ede9fe", "#5b21b6"   # soft purple — no match
            parts.append(
                f"<span style='background:{bg}; color:{fg}; border-radius:4px;"
                f" padding:2px 8px; font-size:0.8rem; margin:2px; display:inline-block;'>{t}</span>"
            )
        return " ".join(parts)

    st.markdown(
        f"""
<div style='display:flex; gap:2rem; flex-wrap:wrap; margin-bottom:0.25rem;'>
  <div style='flex:1; min-width:180px;'>
    <div style='font-size:0.72rem; font-weight:700; letter-spacing:0.07em;
                text-transform:uppercase; color:#6b7280; margin-bottom:4px;'>Coding</div>
    {_tags(coding)}
  </div>
  <div style='flex:2; min-width:220px;'>
    <div style='font-size:0.72rem; font-weight:700; letter-spacing:0.07em;
                text-transform:uppercase; color:#6b7280; margin-bottom:4px;'>Libraries & Frameworks</div>
    {_tags(libs)}
  </div>
  <div style='flex:1.5; min-width:180px;'>
    <div style='font-size:0.72rem; font-weight:700; letter-spacing:0.07em;
                text-transform:uppercase; color:#6b7280; margin-bottom:4px;'>Electrical Engineering</div>
    {_tags(ee)}
  </div>
  <div style='flex:1; min-width:160px;'>
    <div style='font-size:0.72rem; font-weight:700; letter-spacing:0.07em;
                text-transform:uppercase; color:#6b7280; margin-bottom:4px;'>Other</div>
    {_tags(other)}
  </div>
</div>
""",
        unsafe_allow_html=True,
    )
    st.divider()

    # ── 4-column body ──
    c1, c2, c3, c4 = st.columns(4)

    def _bullet_list(col, heading: str, items: list[str]):
        with col:
            st.markdown(f"**{heading}**")
            if not items:
                st.markdown("*None listed*")
                return
            lines = []
            for item in items:
                if _text_matches(item, jd_kw):
                    lines.append(
                        f"<li style='color:#1d4ed8; background:#eff6ff; border-radius:4px;"
                        f" padding:2px 6px; margin:3px 0;'>{item}</li>"
                    )
                else:
                    lines.append(f"<li style='margin:3px 0;'>{item}</li>")
            st.markdown(
                f"<ul style='padding-left:1.2rem; margin:0;'>{''.join(lines)}</ul>",
                unsafe_allow_html=True,
            )

    _bullet_list(c1, "Internships",     r.get("internships", []))
    _bullet_list(c2, "Projects",        r.get("projects", []))
    _bullet_list(c3, "Extracurriculars",r.get("extracurriculars", []))

    _bullet_list(c4, "Key Strengths", r.get("key_strengths", []))
    with c4:
        st.markdown("**Key Gaps**")
        gaps = r.get("key_gaps", [])
        if not gaps:
            st.markdown("*None listed*")
        else:
            lines = [
                f"<li style='color:#dc2626; margin:3px 0;'>{g}</li>"
                for g in gaps
            ]
            st.markdown(
                f"<ul style='padding-left:1.2rem; margin:0;'>{''.join(lines)}</ul>",
                unsafe_allow_html=True,
            )


def build_display_dataframe(results: list[dict]) -> pd.DataFrame:
    rows = []
    for rank, r in enumerate(results, start=1):
        rows.append({
            "Rank": rank,
            "Name": r.get("candidate_name", "Unknown"),
            "Score": r["match_score"],
            "University": r.get("university", "Not listed"),
            "Major": r.get("major", "Not listed"),
        })
    return pd.DataFrame(rows)


def main():
    st.set_page_config(page_title="Resume Screener", layout="wide")

    st.markdown(
        """
<style>
footer { visibility: hidden; }

.main .block-container {
    padding-top: 2rem;
    padding-bottom: 3rem;
    max-width: 1200px;
}

/* Hero banner */
.hero {
    background: linear-gradient(135deg, #5865F2 0%, #7C3AED 100%);
    border-radius: 14px;
    padding: 2rem 2.5rem;
    margin-bottom: 2rem;
    color: white;
}
.hero h1 {
    margin: 0 0 0.3rem 0;
    font-size: 2rem;
    font-weight: 700;
    color: white !important;
}
.hero p {
    margin: 0;
    font-size: 1rem;
    opacity: 0.85;
    color: white !important;
}

/* Section cards */
.upload-card {
    background: white;
    border-radius: 12px;
    padding: 1.25rem 1.5rem 1rem 1.5rem;
    border: 1px solid #e5e7eb;
    box-shadow: 0 1px 6px rgba(0,0,0,0.05);
    margin-bottom: 0.5rem;
}
.section-label {
    font-size: 0.8rem;
    font-weight: 600;
    letter-spacing: 0.07em;
    text-transform: uppercase;
    color: #6b7280;
    margin-bottom: 0.6rem;
}

/* File uploader */
[data-testid="stFileUploader"] {
    background: #fafafa;
    border-radius: 10px;
    padding: 4px;
}

/* Text area */
[data-testid="stTextArea"] textarea {
    border-radius: 8px;
    font-size: 0.875rem;
    border-color: #e5e7eb;
}

/* Primary button */
.stButton > button[kind="primary"] {
    background: linear-gradient(135deg, #5865F2 0%, #7C3AED 100%);
    border: none;
    border-radius: 8px;
    font-weight: 600;
    padding: 0.5rem 2.5rem;
    color: white;
    font-size: 0.95rem;
}
.stButton > button[kind="primary"]:hover {
    background: linear-gradient(135deg, #4752d8 0%, #6929c4 100%);
    border: none;
}

/* Download button */
.stDownloadButton > button {
    border-radius: 8px;
    font-size: 0.85rem;
}

/* Results table */
[data-testid="stDataFrame"] {
    border-radius: 10px;
    overflow: hidden;
    border: 1px solid #e5e7eb;
}

/* Sidebar */
[data-testid="stSidebar"] {
    border-right: 1px solid #e5e7eb;
}
[data-testid="stSidebar"] .stSelectbox label,
[data-testid="stSidebar"] .stNumberInput label {
    font-size: 0.85rem;
    color: #6b7280;
}
</style>
""",
        unsafe_allow_html=True,
    )

    st.markdown(
        """
<div class="hero">
  <h1>Resume Screener</h1>
  <p>Upload a stack of resumes and a job description — we'll read every one and surface your best candidates.</p>
</div>
""",
        unsafe_allow_html=True,
    )

    # ── Sidebar: model selection + API key status ──
    st.sidebar.header("Settings")
    model_label = st.sidebar.selectbox("AI Model", list(MODEL_OPTIONS.keys()))
    provider = MODEL_OPTIONS[model_label]
    num_to_pick = st.sidebar.number_input(
        "How many candidates to pick",
        min_value=1,
        max_value=100,
        value=10,
        step=1,
        help="Number of top candidates to highlight in the results.",
    )

    key_name = API_KEY_MAP[provider]
    api_key = os.environ.get(key_name)
    if not api_key:
        try:
            api_key = st.secrets.get(key_name)
        except Exception:
            api_key = None

    if not api_key:
        st.sidebar.error(f"{key_name} not set.")
        st.error(
            f"{key_name} environment variable is not set. "
            f"Set it in your shell or add it to a .env file alongside app.py."
        )
        st.stop()
    else:
        st.sidebar.success(f"{key_name} loaded.")
        st.sidebar.caption(f"Model: {model_label}")

    # ── Inputs ──
    col1, col2 = st.columns([1, 1])
    with col1:
        uploaded_files = st.file_uploader(
            "Upload Resumes (PDF or DOCX)",
            type=["pdf", "docx"],
            accept_multiple_files=True,
            help="Select one or more PDF or Word resume files.",
        )
        is_combined_pdf = st.checkbox(
            "This is a combined PDF (multiple resumes in one file)",
            value=False,
            help="Check this if you uploaded a single PDF containing multiple resumes. "
                 "The app will auto-detect where each resume starts.",
        )
    with col2:
        st.markdown("**Job Description / Position Prompt**")
        jd_file = st.file_uploader(
            "Upload job description (PDF or DOCX) — optional",
            type=["pdf", "docx"],
            accept_multiple_files=False,
            key="jd_file",
        )
        jd_text_input = st.text_area(
            "Or paste job description here",
            height=220,
            placeholder=(
                "Paste the full job description here, including required skills, "
                "experience level, and responsibilities..."
            ),
        )
        # File upload takes priority; fall back to text area
        if jd_file is not None:
            job_description = extract_text(jd_file) or ""
            if job_description:
                st.success(f"Loaded job description from {jd_file.name}")
            else:
                st.error(f"Could not extract text from {jd_file.name}")
                job_description = jd_text_input
        else:
            job_description = jd_text_input

    inputs_ready = bool(uploaded_files) and bool(job_description.strip())
    run_button = st.button("Screen Resumes", disabled=not inputs_ready, type="primary")

    if not inputs_ready and not run_button:
        if not uploaded_files:
            st.markdown(
                "<div style='text-align:center; padding:3rem 0; color:#9ca3af;'>"
                "<p style='font-size:2rem; margin-bottom:0.5rem;'>📄</p>"
                "<p style='font-size:1rem; font-weight:500; color:#6b7280;'>Upload resumes on the left, paste a job description on the right.</p>"
                "<p style='font-size:0.85rem;'>PDF &amp; DOCX &nbsp;·&nbsp; Batch upload &nbsp;·&nbsp; Combined PDFs supported</p>"
                "</div>",
                unsafe_allow_html=True,
            )
        elif not job_description.strip():
            st.markdown(
                "<div style='text-align:center; padding:2rem 0; color:#9ca3af;'>"
                "<p style='font-size:1.5rem; margin-bottom:0.4rem;'>✏️</p>"
                "<p style='font-size:0.95rem; color:#6b7280;'>Almost there — paste or upload a job description on the right.</p>"
                "</div>",
                unsafe_allow_html=True,
            )

    if run_button:
        # ── Processing ──
        if provider == "claude":
            client = anthropic.Anthropic(api_key=api_key)
        elif provider in ("gemini_31_pro", "gemini_pro", "gemini_flash"):
            client = google_genai.Client(api_key=api_key)
        else:
            client = OpenAI(api_key=api_key)

        # ── Phase 1: Build resume entry list ──
        resume_entries: list[tuple[str, str]] = []
        failures = []

        if is_combined_pdf:
            if len(uploaded_files) != 1 or not uploaded_files[0].name.lower().endswith(".pdf"):
                st.warning(
                    "Combined PDF mode works with a single PDF file. "
                    "Processing each file as a separate resume instead."
                )
                is_combined_pdf = False

        if is_combined_pdf:
            entries = split_combined_pdf(uploaded_files[0])
            if not entries:
                failures.append(f"{uploaded_files[0].name} (could not split combined PDF)")
            else:
                resume_entries = entries
                st.info(f"Split combined PDF into {len(entries)} resume(s).")
        else:
            for uploaded_file in uploaded_files:
                text = extract_text(uploaded_file)
                if text:
                    resume_entries.append((uploaded_file.name, text))
                else:
                    failures.append(f"{uploaded_file.name} (could not extract text)")

        # ── Phase 2: Score each entry ──
        results = []
        total = len(resume_entries)
        progress_bar = st.progress(0, text="Starting...")

        for i, (label, resume_text) in enumerate(resume_entries):
            progress_bar.progress(
                i / total if total > 0 else 0,
                text=f"Analyzing {label} ({i + 1}/{total})...",
            )

            analysis = analyze_resume(client, provider, job_description, resume_text, label)
            if analysis and "_error" not in analysis:
                analysis["filename"] = label
                results.append(analysis)
            else:
                error_msg = analysis.get("_error", "unknown error") if analysis else "no response"
                failures.append(f"{label}: {error_msg}")

        progress_bar.progress(1.0, text="Done.")

        if failures:
            st.warning(
                f"Could not process {len(failures)} file(s):\n"
                + "\n".join(f"- {f}" for f in failures)
            )

        if not results:
            st.error("No resumes were successfully analyzed.")
            return

        results.sort(key=lambda x: x.get("match_score", 0), reverse=True)
        st.session_state["screening_results"] = results
        st.session_state["job_description"] = job_description

    # ── Results (persisted across reruns) ──
    results = st.session_state.get("screening_results")
    if results:
        if len(results) <= num_to_pick:
            top_n = results
            tie_msg = None
        else:
            cutoff_score = results[num_to_pick - 1].get("match_score", 0)
            top_n = [r for r in results if r.get("match_score", 0) >= cutoff_score]
            if len(top_n) > num_to_pick:
                tied_count = sum(1 for r in results if r.get("match_score", 0) == cutoff_score)
                tie_msg = (
                    f"Showing **{len(top_n)} candidates** instead of {num_to_pick} — "
                    f"**{tied_count} candidates are tied at score {cutoff_score}**, so all are included."
                )
            else:
                tie_msg = None

        st.success(f"Screened {len(results)} resume(s) — showing your top {len(top_n)}. Click any row for details.")
        if tie_msg is not None:
            st.info(tie_msg)

        df_display = build_display_dataframe(top_n)
        event = st.dataframe(
            df_display,
            use_container_width=True,
            hide_index=True,
            on_select="rerun",
            selection_mode="single-row",
            column_config={
                "Score": st.column_config.ProgressColumn(
                    "Score",
                    min_value=0,
                    max_value=10,
                    format="%d / 10",
                ),
            },
        )
        selected_rows = event.selection.rows
        if selected_rows:
            jd = st.session_state.get("job_description", "")
            jd_kw = _jd_keywords(jd)
            show_candidate_dialog(top_n[selected_rows[0]], jd_kw, jd)

        df_all = build_dataframe(results)
        csv = df_all.to_csv(index=False)
        st.download_button(
            label=f"Download All {len(results)} Results as CSV",
            data=csv,
            file_name="screening_results.csv",
            mime="text/csv",
        )


if __name__ == "__main__":
    main()
